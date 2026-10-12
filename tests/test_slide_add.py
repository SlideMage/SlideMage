"""Everything the Add slide sheet needs besides widgets: new slide files from the frames, copies and
translations, the deck call for a way of loading a slide, and the slides to pick from."""
from __future__ import annotations

import pytest

from slidemage import alt_blocks, deck_editor, project, slide_add, slide_ids, slide_import, slide_versioning
from slidemage.ai.providers import ChatReply, ProviderError

from conftest import TEMPLATE, use_new_slide_files
from slidemage.project import ProjectContext


@pytest.fixture
def relinked(tmp_path, monkeypatch):
    """The throwaway project; how often the translation links were regenerated."""
    return use_new_slide_files(monkeypatch, tmp_path)


@pytest.fixture
def files(tmp_path, relinked):
    return tmp_path / "slides"


class Translator:
    """A model that answers every request with `reply` (or fails with it, if an exception)."""

    is_local = True
    model = "fake"
    label = "fake"

    def __init__(self, reply):
        self.reply = reply
        self.asked = []

    def chat(self, messages, **kwargs):
        self.asked.append(messages[0]["content"])
        if isinstance(self.reply, Exception):
            raise self.reply
        return ChatReply(text=self.reply)

    def abort(self):
        pass


# -- frames --------------------------------------------------------------------------------

def test_every_frame_has_a_thumbnail_and_a_body_with_the_title(ctx):
    assert slide_add.FRAMES[0] == slide_add.STANDARD_FRAME
    for frame in slide_add.FRAMES:
        assert slide_add.frame_thumbnail(frame).is_file(), frame
    for frame in slide_add.FRAMES[1:]:
        body = slide_add.frame_body(ctx, frame, "Growth & R_D")
        assert "{Growth \\& R\\_D}" in body or "Growth \\& R\\_D" in body, frame
        assert "\n" not in body.replace("\r\n", ""), frame
    for frame in ("standardframe", "figureframe", "sideimageframe", "highlightframe", "sectionframe",
                  "titleframe"):
        assert alt_blocks.frame_type(slide_add.frame_body(ctx, frame, "Hello")) == frame


def test_the_standard_frame_is_the_projects_template(ctx, files):
    assert slide_add.frame_body(ctx, "standardframe", "") == TEMPLATE
    assert slide_add.frame_body(ctx, "standardframe", " My title ") == TEMPLATE.replace("{Title}", "{My title}")


@pytest.mark.parametrize("custom_template", [None, TEMPLATE + "% Custom template\r\n"])
def test_empty_project_can_add_a_standard_slide(tmp_path, monkeypatch, custom_template):
    target = tmp_path / "empty-project"
    project.init_project(target, copy_example=False)
    relative = ".scripts/R_Slidemaster/inst/templates/slide-template.tex"
    template = target / relative
    ctx = ProjectContext.at(target)
    expected = (project.template_dir("example") / relative).read_bytes().decode("utf-8")
    if custom_template is not None:
        template.parent.mkdir(parents=True)
        template.write_bytes(custom_template.encode("utf-8"))
        expected = custom_template

    assert slide_ids.read_template(ctx) == expected
    created = slide_add.create_new(ctx, "First", "EN")
    assert created.path.read_bytes().decode("utf-8") == slide_ids.build_new_slide_header() + expected
    assert template.exists() == (custom_template is not None)


@pytest.mark.parametrize("bundled_available", [False, True])
def test_missing_project_and_bundled_templates_raise(tmp_path, monkeypatch, bundled_available):
    ctx = ProjectContext.at(tmp_path)     # no template in it
    monkeypatch.setattr(project, "template_dir", lambda _name: tmp_path if bundled_available else None)

    with pytest.raises(FileNotFoundError, match="No slide template found"):
        slide_ids.read_template(ctx)


# -- new files -------------------------------------------------------------------------------

def test_a_new_slide_gets_a_fresh_id_a_header_and_the_frames_text(ctx, files):
    created = slide_add.create_new(ctx, "My-Slide", "DE", "quoteframe", "Hello")
    bare, language, slug = created.id_slug.split("-", 2)
    assert (language, slug) == ("DE", "My-Slide") and created.path == files / f"{created.id_slug}-current.tex"
    text = created.path.read_bytes().decode("utf-8")
    assert text.startswith("%% Included in decks: (none)\r\n%%\r\n\r\n")
    assert "\\quoteframe{" in text and "\\SMAgendaTitle{Hello}" in text
    assert f'"{bare}"' in (files.parent / "slide_ids.csv").read_text()
    with pytest.raises(ValueError):
        slide_add.create_new(ctx, "Bad", "FR")


def test_a_duplicate_keeps_the_language_and_does_not_link(ctx, files, relinked):
    (files / "abc-DE-Demo-current.tex").write_bytes(
        b"%% Included in decks: talk\r\n%%\r\n\r\n\\begin{standardframe}{Demo}\r\nText\r\n\\end{standardframe}\r\n")
    created = slide_add.create_copy(ctx, "abc-DE-Demo", "Demo-Copy")
    assert created.id_slug.endswith("-DE-Demo-Copy")
    text = created.path.read_bytes().decode("utf-8")
    assert text == "%% Included in decks: (none)\r\n%%\r\n\r\n\\begin{standardframe}{Demo}\r\nText\r\n\\end{standardframe}\r\n"
    assert "translation" not in (files / "abc-DE-Demo-current.tex").read_text() and relinked == []
    legacy = files / "old-slide-current.tex"
    legacy.write_bytes(b"body")
    assert slide_add.create_copy(ctx, "old-slide", "Copy").id_slug.endswith("-EN-Copy")   # never left untagged


def test_a_translation_is_linked_both_ways(ctx, files, relinked):
    source = files / "abc-EN-Demo-current.tex"
    source.write_bytes(b"%% Included in decks: talk\r\n%%\r\n\r\nText\r\n")
    created = slide_add.create_copy(ctx, "abc-EN-Demo", "Demo", "DE")
    assert created.path.read_bytes().decode("utf-8").startswith(
        "%% Included in decks: (none)\r\n%%\r\n%% Translation of: abc-EN-Demo\r\n%%\r\n\r\nText")
    assert f"%% DE translation: {created.id_slug}" in source.read_text()
    assert relinked == [True]
    translated = slide_add.create_copy(ctx, "abc-EN-Demo", "Demo-2", "DE", body="Text auf Deutsch\r\n")
    assert translated.path.read_bytes().decode("utf-8").endswith("\r\n\r\nText auf Deutsch\r\n")


@pytest.mark.parametrize("language", [None, "DE"])
@pytest.mark.parametrize("supplied_body", [False, True])
def test_copies_start_untracked_like_new_slides_without_changing_the_source_body(ctx, files, language, supplied_body):
    source = files / "abc-EN-Demo-current.tex"
    source_body = "\\setSlideVersion{1-1-0}\r\nText\r\n"
    source_bytes = (slide_ids.build_new_slide_header() + source_body).encode("utf-8")
    source.write_bytes(source_bytes)
    new = slide_add.create_new(ctx, "New", "EN")
    copied = slide_add.create_copy(ctx, "abc-EN-Demo", "Copy", language,
                                   body=source_body if supplied_body else None)

    assert slide_versioning.parse_version(new.path.read_text()) is None
    assert slide_versioning.parse_version(copied.path.read_text()) is None
    assert slide_ids.slide_body(copied.path) == "Text\r\n"
    assert slide_ids.slide_body(source) == source_body
    assert slide_versioning.parse_version(source.read_text()) == (1, 1, 0)
    if language is None:
        assert source.read_bytes() == source_bytes


def test_the_languages_a_slide_can_be_translated_into(files):
    assert slide_add.other_languages("abc-EN-Demo") == ["DE", "mult"]
    assert slide_add.other_languages("abc-DE-Demo") == ["EN", "mult"]
    assert slide_add.other_languages("abc-Demo") == ["EN", "DE", "mult"]
    assert slide_add.can_machine_translate("abc-EN-Demo", "DE")
    assert not slide_add.can_machine_translate("abc-DE-Demo", "EN")
    assert not slide_add.can_machine_translate("abc-EN-Demo", "mult")
    assert slide_add.slug_of("abc-EN-Demo-Two") == "Demo-Two" and slide_add.slug_of("abc-Demo") == "Demo"


def test_slugs():
    assert slide_import.slug_from_title("Growth of R&D spend (50% YoY)", "fallback") == \
        "Growth-of-R-D-spend-50-YoY"
    assert slide_import.slug_from_title("", "Imported-Slide-1") == "Imported-Slide-1"
    assert slide_import.slug_from_title("!!!", "Imported-Slide-2") == "Imported-Slide-2"
    assert slide_import.slug_from_title("???", "") == ""
    assert all(slide_add.SLUG_RE.match(s) for s in ("MM-Theorem", "a", "A1-b2"))
    assert not any(slide_add.SLUG_RE.match(s) for s in ("", "-a", "a-", "a--b", "a b", "ä"))


# -- translating -----------------------------------------------------------------------------

def test_the_body_is_translated_as_one_block_and_keeps_its_line_ends(ctx):
    model = Translator("Hallo\nWelt")
    text, error = slide_add.translate_body(ctx, model, "Hello\r\nWorld\r\n")
    assert (text, error) == ("Hallo\r\nWelt", None) and len(model.asked) == 1 and "Hello\r\nWorld" in model.asked[0]
    assert slide_add.translate_body(ctx, Translator("Hallo\r\nWelt"), "Hello\nWorld")[0] == "Hallo\nWelt"


def test_a_failed_translation_is_reported_not_raised(ctx):
    text, error = slide_add.translate_body(ctx, Translator(ProviderError("rate limited")), "Hello")
    assert (text, error) == ("", "rate limited")


# -- the deck call -----------------------------------------------------------------------------

def test_the_call_follows_the_load_and_the_views():
    views = deck_editor.DEFAULT_CONDITIONS
    assert slide_add.call_spec("live", views) == slide_add.CallSpec("slide")
    assert slide_add.call_spec("live", ["after-course", "instructor"]) == slide_add.CallSpec(
        "conditionalslide", conditions=["instructor", "after-course"])
    assert slide_add.call_spec("frozen", [], "1-0-0") == slide_add.CallSpec("slideversion", "1-0-0")
    assert slide_add.call_spec("step", ["during-course"], steps=" 1, 2 ") == slide_add.CallSpec(
        "slidestep", conditions=["during-course"], steps="1,2")
    assert slide_add.call_spec("step", views).steps == "1"
    for bad in (lambda: slide_add.call_spec("live", []), lambda: slide_add.call_spec("step", views, steps="x"),
                lambda: slide_add.call_spec("frozen", views)):
        with pytest.raises(ValueError):
            bad()
    spec = slide_add.call_spec("step", views, steps="2-")
    assert deck_editor.format_call(spec.macro, "abc-EN-Demo", spec.version, spec.conditions, spec.steps) == \
        "\\slidestep{abc-EN-Demo}[instructor,during-course,after-course]{2-}"


# -- the slides to pick from -----------------------------------------------------------------------

def test_the_project_slides_with_title_and_language(ctx, files):
    (files / "bbb-DE-Zwei-current.tex").write_text("\\begin{standardframe}{Zwei}\nx\n\\end{standardframe}\n")
    (files / "aaa-EN-One-current.tex").write_text("\\begin{standardframe}{One}\nx\n\\end{standardframe}\n")
    (files / "aaa-EN-One-v1-0-0.tex").write_text("frozen")
    assert slide_add.project_slides(ctx) == [
        slide_add.ProjectSlide("aaa-EN-One", "One", "[EN]"), slide_add.ProjectSlide("bbb-DE-Zwei", "Zwei", "[DE]")]


# -- presentation import -------------------------------------------------------------

_FAKE_SVM = b"VCLMTF\x01\x001\x00\x00\x00"


def _make_test_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Growth of R&D spend (50% YoY)"
    tf = slide.placeholders[1].text_frame
    tf.text = "Top level point #1"
    p2 = tf.add_paragraph(); p2.text = "Nested detail A"; p2.level = 1
    p3 = tf.add_paragraph(); p3.text = "Nested detail B"; p3.level = 1
    p4 = tf.add_paragraph(); p4.text = "Second top-level point with $5 cost"; p4.level = 0

    img_path = path.parent / "red.png"
    from PIL import Image
    Image.new("RGB", (200, 100), color=(255, 0, 0)).save(img_path)
    slide.shapes.add_picture(str(img_path), Inches(1), Inches(4))

    image_only = prs.slides.add_slide(prs.slide_layouts[6])
    image_only.shapes.add_picture(str(img_path), Inches(1), Inches(1))

    prs.slides.add_slide(prs.slide_layouts[6])  # blank
    prs.save(str(path))


def _make_test_odp(path):
    from odf.draw import Frame, Image, Page, TextBox
    from odf.namespaces import PRESENTATIONNS
    from odf.opendocument import OpenDocumentPresentation
    from odf.style import MasterPage, PageLayout
    from odf.text import P, List, ListItem

    doc = OpenDocumentPresentation()
    layout = PageLayout(name="PM1")
    doc.automaticstyles.addElement(layout)
    master = MasterPage(name="Default", pagelayoutname=layout)
    doc.masterstyles.addElement(master)

    def make_page(name):
        p = Page(name=name, masterpagename=master)
        doc.presentation.addElement(p)
        return p

    page1 = make_page("Slide1")
    title_frame = Frame(width="25cm", height="3cm", x="1cm", y="1cm")
    title_frame.setAttrNS(PRESENTATIONNS, "class", "title")
    tb = TextBox(); tb.addElement(P(text="Growth of R&D spend (50% YoY)"))
    title_frame.addElement(tb); page1.addElement(title_frame)

    content_frame = Frame(width="20cm", height="10cm", x="1cm", y="4cm")
    content_frame.setAttrNS(PRESENTATIONNS, "class", "outline")
    tb2 = TextBox()
    top_level = List()
    top_item = ListItem(); top_item.addElement(P(text="Top level point #1"))
    nested = List()
    li1 = ListItem(); li1.addElement(P(text="Nested detail A")); nested.addElement(li1)
    li2 = ListItem(); li2.addElement(P(text="Nested detail B")); nested.addElement(li2)
    top_item.addElement(nested)
    top_level.addElement(top_item)
    second_item = ListItem()
    second_item.addElement(P(text="Second top-level point with $5 cost"))
    top_level.addElement(second_item)
    tb2.addElement(top_level)
    content_frame.addElement(tb2); page1.addElement(content_frame)

    img_path = path.parent / "red.png"
    from PIL import Image as PILImage
    PILImage.new("RGB", (200, 100), color=(255, 0, 0)).save(img_path)
    img_href = doc.addPicture(str(img_path))
    img_frame = Frame(width="5cm", height="3cm", x="2cm", y="15cm")
    img_frame.addElement(Image(href=img_href)); page1.addElement(img_frame)

    page2 = make_page("Slide2")
    img_frame2 = Frame(width="10cm", height="6cm", x="2cm", y="2cm")
    img_frame2.addElement(Image(href=img_href)); page2.addElement(img_frame2)

    make_page("Slide3")  # blank
    doc.save(str(path))


@pytest.fixture(params=["pptx", "odp"])
def sample_presentation(request, tmp_path):
    from slidemage import slide_import

    if request.param == "pptx":
        path = tmp_path / "sample.pptx"
        _make_test_pptx(path)
        return path, slide_import.extract_pptx
    path = tmp_path / "sample.odp"
    _make_test_odp(path)
    return path, slide_import.extract_odp


def test_escape_latex_special_characters():
    from slidemage.slide_import import escape_latex

    assert escape_latex("50% off & $5 #1 a_b {x} ~y ^z \\end") == (
        r"50\% off \& \$5 \#1 a\_b \{x\} \textasciitilde{}y "
        r"\textasciicircum{}z \textbackslash{}end"
    )


def _odp_with_one_page(tmp_path, build_page):
    from odf.draw import Page
    from odf.opendocument import OpenDocumentPresentation
    from odf.style import MasterPage, PageLayout

    doc = OpenDocumentPresentation()
    layout = PageLayout(name="PM1")
    doc.automaticstyles.addElement(layout)
    master = MasterPage(name="Default", pagelayoutname=layout)
    doc.masterstyles.addElement(master)
    page = Page(name="Slide1", masterpagename=master)
    doc.presentation.addElement(page)
    build_page(doc, page)
    path = tmp_path / "sample.odp"
    doc.save(str(path))
    return path


def test_extract_odp_skips_starview_metafile_masquerading_as_png(tmp_path):
    # LibreOffice's VCLMTF format, labeled image/png in the manifest, used
    # to be written as "imported-...png" and broke the PDF build (libpng:
    # "Not a PNG file").
    from odf.draw import Frame, Image
    from slidemage import slide_import

    def build(doc, page):
        href = doc.addPicture("fake.png", mediatype="image/png", content=_FAKE_SVM)
        frame = Frame(width="5cm", height="3cm", x="1cm", y="1cm")
        frame.addElement(Image(href=href))
        page.addElement(frame)

    slides = slide_import.extract_odp(_odp_with_one_page(tmp_path, build))
    assert all(not slide.images for slide in slides)


def test_extract_odp_imports_table_instead_of_its_preview(tmp_path):
    # A table frame carries both the <table:table> and a TablePreview .svm
    # picture of it; the table's own cells are the content.
    from odf.draw import Frame, Image, TextBox
    from odf.namespaces import PRESENTATIONNS
    from odf.table import CoveredTableCell, Table, TableCell, TableColumn, TableRow
    from odf.text import P
    from slidemage import slide_import

    def build(doc, page):
        table = Table()
        for _ in range(3):
            table.addElement(TableColumn())
        header = TableRow()
        c1 = TableCell(); c1.addElement(P(text="Date")); header.addElement(c1)
        c2 = TableCell(numbercolumnsspanned=2)
        c2.addElement(P(text="Topic")); c2.addElement(P(text="Lecturer"))
        header.addElement(c2); header.addElement(CoveredTableCell())
        table.addElement(header)
        body = TableRow()
        for text in ["7.10.", "Intro & SQL", ""]:
            cell = TableCell()
            if text:
                cell.addElement(P(text=text))
            body.addElement(cell)
        table.addElement(body)

        table_frame = Frame(width="20cm", height="10cm", x="1cm", y="2cm")
        table_frame.addElement(table)
        href = doc.addPicture("Pictures/TablePreview1.svm", mediatype="image/png", content=_FAKE_SVM)
        table_frame.addElement(Image(href=href))
        page.addElement(table_frame)

        number_frame = Frame(width="5cm", height="1cm", x="1cm", y="18cm")
        number_frame.setAttrNS(PRESENTATIONNS, "class", "page-number")
        tb = TextBox(); tb.addElement(P(text="<Foliennummer>"))
        number_frame.addElement(tb); page.addElement(number_frame)

    [slide] = slide_import.extract_odp(_odp_with_one_page(tmp_path, build))
    assert slide.images == []
    assert slide.bullets == []
    assert [[(c.paragraphs, c.colspan) for c in row] for row in slide.tables[0]] == [
        [(["Date"], 1), (["Topic", "Lecturer"], 2)],
        [(["7.10."], 1), (["Intro & SQL"], 1), ([], 1)],
    ]

    layout, tex = slide_import.slide_content_to_tex(slide, [])
    assert layout == "standardframe"
    assert r"\begin{gothamtable}[gap=12pt]{lll}" in tex
    assert (
        r"Date & \multicolumn{2}{l@{}}{\begin{tabular}[t]{@{}l@{}}Topic \\ Lecturer\end{tabular}} \\"
        in tex
    )
    assert r"7.10. & Intro \& SQL &  \\" in tex


def test_item_text_starting_with_overlay_bracket_is_protected():
    from slidemage.slide_import import Bullet, _bullets_to_itemize

    tex = _bullets_to_itemize([Bullet("<Foliennummer>"), Bullet("[optional] note")])
    assert r"\item {}<Foliennummer>" in tex
    assert r"\item {}[optional] note" in tex



@pytest.mark.parametrize("children,expected", [
    ([("deep", 2)], "\t\t\t\\item deep\n"),
    ([("deep A", 3), ("deep B", 3)], "\t\t\t\\item deep A\n\t\t\t\\item deep B\n"),
], ids=["missing-level", "repeated-depth-jump-regression"])
def test_bullets_to_itemize_nesting_and_jumps(children, expected):
    from slidemage.slide_import import Bullet, _bullets_to_itemize

    tex = _bullets_to_itemize([
        Bullet("top", level=0),
        *(Bullet(text, level=level) for text, level in children),
        Bullet("back to top", level=0),
    ])
    assert tex == (
        "\t\\begin{itemize}\n"
        "\t\t\\item top\n"
        "\t\t\\begin{itemize}\n"
        + expected +
        "\t\t\\end{itemize}\n"
        "\t\t\\item back to top\n"
        "\t\\end{itemize}"
    )


def test_extract_matches_across_pptx_and_odp(sample_presentation):
    path, extractor = sample_presentation
    slides = extractor(path)

    assert len(slides) == 3
    text_slide, image_only_slide, blank_slide = slides

    assert text_slide.title == "Growth of R&D spend (50% YoY)"
    assert [(b.text, b.level) for b in text_slide.bullets] == [
        ("Top level point #1", 0),
        ("Nested detail A", 1),
        ("Nested detail B", 1),
        ("Second top-level point with $5 cost", 0),
    ]
    assert len(text_slide.images) == 1
    assert text_slide.images[0][0][:4] == b"\x89PNG"
    assert text_slide.has_text and not text_slide.is_blank

    assert not image_only_slide.has_text
    assert len(image_only_slide.images) == 1
    assert not image_only_slide.is_blank

    assert blank_slide.is_blank


def test_slide_content_to_tex_standardframe_vs_figureframe():
    from slidemage.slide_import import Bullet, SlideContent, slide_content_to_tex

    text_content = SlideContent(title="A title", bullets=[Bullet("point", level=0)])
    layout, tex = slide_content_to_tex(text_content, [])
    assert layout == "standardframe"
    assert tex.startswith(r"\begin{standardframe}{A title}")
    assert r"\item point" in tex

    image_content = SlideContent(images=[(b"x", ".png")])
    layout, tex = slide_content_to_tex(image_content, ["media/x.png"])
    assert layout == "figureframe"
    assert tex.startswith(r"\begin{figureframe}{Untitled}")
    assert r"\includegraphics[height=0.72\textheight]{media/x.png}" in tex


def test_import_presentation_writes_files_and_appends_deck_calls(sample_presentation, tmp_path, monkeypatch):
    from slidemage import deck_scan, slide_import

    path, _extractor = sample_presentation
    project_root = tmp_path / "project"
    (project_root / "slides").mkdir(parents=True)
    (project_root / "decks").mkdir()
    ctx = ProjectContext.at(project_root)

    deck_text = "\\documentclass{beamer}\n\\begin{document}\n\\end{document}\n"
    deck_path = project_root / "decks" / "test.tex"
    deck_path.write_text(deck_text, encoding="utf-8", newline="")
    deck = deck_scan.Deck(
        name="test", path=deck_path, archived=False, text=deck_text, supports_version=False,
        uses_biblatex=False, preamble_end=None, chapters=[], output_folders={}, tags=[],
    )

    summary = slide_import.import_presentation(ctx, path, deck, "EN")

    assert len(summary.imported) == 2
    assert summary.skipped_blank == 1
    assert [s.layout for s in summary.imported] == ["standardframe", "figureframe"]

    new_deck_text = deck_path.read_bytes().decode("utf-8")
    for slide in summary.imported:
        assert f"\\slide{{{slide.id_slug}}}" in new_deck_text
        slide_file = project_root / "slides" / f"{slide.id_slug}-current.tex"
        assert slide_file.is_file()

    csv_text = (project_root / "slide_ids.csv").read_bytes().decode("utf-8")
    for slide in summary.imported:
        expected_id = slide.id_slug.split("-", 1)[0]
        assert f'"{expected_id}"' in csv_text

    media_files = list((project_root / "media").glob("*.png"))
    assert len(media_files) == 2
