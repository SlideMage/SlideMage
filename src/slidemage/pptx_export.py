"""Exports a compiled deck PDF as a PowerPoint file with one full-bleed
image per PDF page (2026-08-27, direct instruction) -- lets PowerPoint's
own recording/presenter tools be used with a SlideMaster deck, which
PowerPoint can't open directly (it's a PDF, not a native .pptx).

Every beamer overlay/reveal step of one \\slide{} call is stacked onto a
single PowerPoint slide instead of getting its own (2026-08-27, direct
instruction), so an annotation drawn during a live recording stays
visible across that slide's own later reveal steps -- PowerPoint's own
ink annotations paint in a layer above everything on the slide, so they
survive regardless of which image beneath them is currently shown.
Reveal steps are grouped by their shared footer "slide id: XXX" text
(every beamer overlay of one \\slide{} call carries the identical id --
see slide_snapshot.py's slide_id_of, used here), each stacked image
full-bleed like a single-page slide would
be, ordered so the *first* reveal step ends up on top (highest z-order --
document order in <p:spTree> is what determines it, later added is drawn
over earlier) and the *last*/fullest reveal step ends up on the bottom.
Every image but the bottommost gets a click-triggered "Fade" ("Verblassen",
2026-08-28, direct instruction -- was "Disappear"/"Verschwinden" until
then) exit animation, so each click reveals whatever's stacked beneath
it -- since beamer overlays are themselves cumulative (image N always
contains everything image N-1 had, plus more), stripping away the topmost
layer one click at a time reproduces the same reveal sequence PowerPoint's
own way.

Uses the python-pptx package (direct instruction, weighed against
hand-building the OOXML package by hand -- see requirements.txt's own
comment on why): a deliberate exception to this project's stdlib-first
policy. python-pptx itself has no animation API at all though (confirmed
directly via web research 2026-08-27, including a blog post whose own
author called every attempt at this "stalled on edge cases") -- the
<p:timing> tree below is raw OOXML, hand-assembled and appended straight
onto each slide's own XML element. Verified directly (2026-08-27/28) via
computer-use: built the real "Verschwinden"/Disappear and
"Verblassen"/Fade effects in a real PowerPoint installation, saved, and
read their own generated XML back -- see _fade_click_effect()'s own
docstring for what that ground truth actually looks like.
"""
from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Callable, Optional

import pymupdf as fitz
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from .slide_snapshot import slide_id_of

LogCallback = Callable[[str], None]

# Standard PowerPoint 16:9 widescreen slide size, in EMU (914400 EMU/inch)
# -- 13.333in x 7.5in. Matches this project's own beamer aspectratio=169
# decks exactly (confirmed directly against a real compiled PDF: page
# ratio 1.7778), so every rendered page fills the slide with no
# letterboxing on either axis.
SLIDE_WIDTH_EMU = 12192000
SLIDE_HEIGHT_EMU = 6858000

# Rendered page width in pixels -- a fixed target regardless of a given
# deck's own PDF page size in points, so output resolution stays
# consistent across decks. ~2560px wide is comfortably sharp for a full-
# screen recording without producing an unreasonably large .pptx for a
# multi-hour lecture deck.
_TARGET_WIDTH_PX = 2560

# NOTE: "openxmlformats.org", not "openxml.org" -- a typo here (missing
# "formats", 2026-08-27/28, found directly by comparing this constant
# against a real PowerPoint-saved file's own root xmlns:p) silently
# breaks every animation this module generates: PowerPoint accepts the
# resulting .pptx with no repair prompt (still schema-valid XML) but
# treats <p:timing> as belonging to an unrecognized namespace and drops
# it entirely -- Animation Pane shows nothing, no error anywhere.
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def render_pdf_pages_grouped(pdf_path: Path, out_dir: Path) -> list[list[Path]]:
    """PNGs for every PDF page (written into `out_dir`, created if
    needed), grouped by consecutive pages sharing the same footer slide
    id -- one group per beamer slide, reveal steps in order within each
    group, groups in page order. A page whose footer carries no
    recognizable slide id (shouldn't happen for a real compiled deck, but
    not assumed) is always its own single-page group, never merged with a
    neighbor by coincidence."""
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    try:
        digits = max(4, len(str(doc.page_count)))
        groups: list[list[Path]] = []
        current_id: Optional[str] = None
        for index in range(doc.page_count):
            page = doc[index]
            slide_id = slide_id_of(page)
            zoom = _TARGET_WIDTH_PX / page.rect.width
            pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
            png_path = out_dir / f"page-{index + 1:0{digits}d}.png"
            pixmap.save(png_path)
            if slide_id is not None and slide_id == current_id:
                groups[-1].append(png_path)
            else:
                groups.append([png_path])
                current_id = slide_id
        return groups
    finally:
        doc.close()


# The "Verblassen"/Fade exit effect's own interpolated duration, in
# milliseconds -- matches PowerPoint's own default when applying "Fade"
# from the ribbon gallery by hand (confirmed directly, 2026-08-27).
_FADE_DURATION_MS = 500


def _fade_click_effect(effect_id: int, shape_id: int) -> str:
    """One click-triggered "Fade" ("Verblassen") exit effect, targeting
    `shape_id`: an interpolated opacity fade (<p:animEffect
    transition="out" filter="fade">, over _FADE_DURATION_MS) followed by
    an instant style.visibility=hidden toggle (<p:set>) once the fade
    finishes, so the shape is actually removed from hit-testing/rendering
    rather than just left transparent. Ground truth for this whole
    structure (2026-08-27/28, confirmed directly via computer-use): built
    the real *entrance* "Verblassen"/Fade effect in a real PowerPoint
    installation (Animations tab -> ribbon gallery), saved, and read its
    own generated XML back -- presetID="10", the three-level <p:cTn>
    click-wait nesting (same shape a "Verschwinden"/Disappear effect uses,
    see this module's own git history), and a sibling <p:animEffect> next
    to the <p:set> are all exactly what PowerPoint itself wrote. The exit
    (as opposed to entrance) variant used here --
    presetClass="exit"/transition="out"/val="hidden", <p:animEffect>
    before <p:set>, and the <p:set> delayed by _FADE_DURATION_MS instead
    of firing at delay="0" -- was hand-derived by mirroring that captured
    entrance structure and then verified by hand-patching a saved
    PowerPoint file's own XML with exactly this candidate and confirming
    directly, via computer-use, that PowerPoint's Animation Pane
    recognizes it as a normal "Verblassen" exit effect (correct red
    Exit-category gallery highlight, correct Start/Dauer fields) and that
    running the slideshow and clicking actually fades the shape out and
    hides it -- not merely that PowerPoint opens the file without
    prompting to repair it, which on its own only proves the XML is
    schema-valid, not that it behaves as intended (this is exactly how
    the very first, wrongly-nested version of this module's Disappear
    effect failed: no repair prompt, empty Animation Pane). `effect_id`
    must be unique within this slide's own timing tree; four more ids are
    consumed after it (the caller must space successive calls at least 5
    apart)."""
    return f"""<p:par xmlns:p="{_P_NS}">
  <p:cTn id="{effect_id}" fill="hold">
    <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
    <p:childTnLst>
      <p:par>
        <p:cTn id="{effect_id + 1}" fill="hold">
          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
          <p:childTnLst>
            <p:par>
              <p:cTn id="{effect_id + 2}" presetID="10" presetClass="exit" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                <p:childTnLst>
                  <p:animEffect transition="out" filter="fade">
                    <p:cBhvr>
                      <p:cTn id="{effect_id + 4}" dur="{_FADE_DURATION_MS}"/>
                      <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                    </p:cBhvr>
                  </p:animEffect>
                  <p:set>
                    <p:cBhvr>
                      <p:cTn id="{effect_id + 3}" dur="1" fill="hold">
                        <p:stCondLst><p:cond delay="{_FADE_DURATION_MS}"/></p:stCondLst>
                      </p:cTn>
                      <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                      <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                    </p:cBhvr>
                    <p:to><p:strVal val="hidden"/></p:to>
                  </p:set>
                </p:childTnLst>
              </p:cTn>
            </p:par>
          </p:childTnLst>
        </p:cTn>
      </p:par>
    </p:childTnLst>
  </p:cTn>
</p:par>"""


def _build_timing_xml(shape_ids_topmost_first: list[int]) -> str:
    """The full <p:timing> tree for one slide: every shape but the last
    in `shape_ids_topmost_first` gets its own sequential click-triggered
    Fade effect, in the order given (topmost/first-reveal shape fades out
    first). The tmRoot -> seq(mainSeq) -> childTnLst wrapper below, and
    the trailing <p:bldLst> (2026-08-27, confirmed directly -- an earlier
    version of this function omitted it entirely, which alone may have
    been enough to keep the Animation Pane empty even before the nesting
    bug above was also fixed) are both exactly what real PowerPoint itself
    writes for any number of sequential click effects on one slide."""
    effects = shape_ids_topmost_first[:-1]  # the bottommost shape never fades out
    next_id = 3
    par_blocks: list[str] = []
    bld_entries: list[str] = []
    for shape_id in effects:
        par_blocks.append(_fade_click_effect(next_id, shape_id))
        bld_entries.append(f'<p:bldP spid="{shape_id}" grpId="0" animBg="1"/>')
        next_id += 5  # this call's own id, +1, +2, +3, +4
    child_tn_lst = "\n".join(par_blocks)
    bld_lst = "".join(bld_entries)
    return f"""<p:timing xmlns:p="{_P_NS}">
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                {child_tn_lst}
              </p:childTnLst>
            </p:cTn>
            <p:prevCondLst>
              <p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:prevCondLst>
            <p:nextCondLst>
              <p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:nextCondLst>
          </p:seq>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
  <p:bldLst>{bld_lst}</p:bldLst>
</p:timing>"""


def build_pptx(page_groups: list[list[Path]], out_path: Path) -> None:
    """A .pptx with one slide per entry in `page_groups`. A single-image
    group is just that image, full-bleed, no animation -- the common
    case for a slide with no beamer overlays. A multi-image group stacks
    every image full-bleed (first reveal on top) with click-triggered
    Fade effects peeling them away in reveal order -- see this module's
    own docstring."""
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(SLIDE_HEIGHT_EMU)
    # By name, not a hardcoded layout index -- more robust against a
    # different python-pptx version ever shipping its bundled default
    # template with a different layout order, though index 6 ("Blank") is
    # itself a longstanding, documented convention. Blank specifically:
    # no inherited placeholders to work around, since every slide here is
    # just full-bleed pictures this function adds itself.
    blank_layout = next(
        (layout for layout in prs.slide_layouts if layout.name == "Blank"),
        prs.slide_layouts[6],
    )
    for group in page_groups:
        slide = prs.slides.add_slide(blank_layout)
        # Added bottommost (fullest reveal) first, topmost (first reveal)
        # last -- <p:spTree> document order is z-order, later added draws
        # over earlier, so this ordering alone is what makes the first
        # reveal step visible at rest.
        shape_ids_topmost_first: list[int] = []
        for image_path in reversed(group):
            picture = slide.shapes.add_picture(
                str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height,
            )
            shape_ids_topmost_first.insert(0, picture.shape_id)
        if len(group) > 1:
            timing_xml = _build_timing_xml(shape_ids_topmost_first)
            timing_element = etree.fromstring(timing_xml.encode("utf-8"))
            slide._element.append(timing_element)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))


def export_deck_to_pptx(
    pdf_path: Path, out_path: Path, log: Optional[LogCallback] = None,
) -> None:
    """Renders every page of `pdf_path`, groups them by beamer slide, and
    assembles them into `out_path` as a .pptx -- see this module's own
    docstring. Raises on failure (a corrupt/unreadable PDF, a python-pptx
    error) -- the caller decides how to surface that; nothing here is
    swallowed."""
    log = log or (lambda line: None)
    with tempfile.TemporaryDirectory(prefix="slidemage_pptx_export_") as tmp:
        tmp_dir = Path(tmp)
        log(f"    Rendering pages from {pdf_path.name}...")
        groups = render_pdf_pages_grouped(pdf_path, tmp_dir)
        page_count = sum(len(g) for g in groups)
        multi = sum(1 for g in groups if len(g) > 1)
        log(f"    Rendered {page_count} page(s) into {len(groups)} slide(s) "
           f"({multi} with reveal-step animations); building {out_path.name}...")
        build_pptx(groups, out_path)
